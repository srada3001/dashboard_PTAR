"""Reads the Excel tables (ListObjects) of a workbook.

openpyxl only exposes tables in normal mode, which fails on these workbooks (pivot-table caches),
so the tables are located in the file's XML and then read in read-only mode (CLAUDE.md). The
workbook is read twice in step, once for the values and once for the formulas: a cell holding a
formula is computed, not measured, and is never imported.
"""

import hashlib
import posixpath
import re
import zipfile
from dataclasses import dataclass
from xml.etree import ElementTree

import openpyxl
from openpyxl.utils import range_boundaries

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


@dataclass(frozen=True)
class ExcelTable:
    name: str
    sheet: str
    ref: str  # "A7:AW38434", header row included
    totals_rows: int = 0


@dataclass(frozen=True)
class Cell:
    value: object
    formula: bool


@dataclass
class TableData:
    table: ExcelTable
    headers: list  # raw header values
    rows: list  # (excel row number, [Cell, ...])


def _target(base, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), target))


def _relationships(archive, part):
    rels_path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rels_path not in archive.namelist():
        return {}
    root = ElementTree.fromstring(archive.read(rels_path))
    return {rel.get("Id"): _target(part, rel.get("Target")) for rel in root.findall("rel:Relationship", NS)}


def find_tables(path):
    """{table name: ExcelTable} for every Excel table in the workbook."""
    tables = {}
    with zipfile.ZipFile(path) as archive:
        workbook = "xl/workbook.xml"
        sheet_parts = _relationships(archive, workbook)
        root = ElementTree.fromstring(archive.read(workbook))
        for sheet in root.findall("main:sheets/main:sheet", NS):
            part = sheet_parts.get(sheet.get(R_ID))
            if not part:
                continue
            for target in _relationships(archive, part).values():
                if not re.search(r"/tables/[^/]+\.xml$", target):
                    continue
                table = ElementTree.fromstring(archive.read(target))
                name = table.get("name") or table.get("displayName")
                tables[name] = ExcelTable(
                    name=name, sheet=sheet.get("name"), ref=table.get("ref"),
                    totals_rows=int(table.get("totalsRowCount") or 0),
                )
    return tables


def read_table(path, table):
    """The header and data rows of an Excel table, each cell with its value and whether it is a formula."""
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    max_row -= table.totals_rows
    values_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        bounds = {"min_row": min_row, "max_row": max_row, "min_col": min_col, "max_col": max_col, "values_only": True}
        values = values_book[table.sheet].iter_rows(**bounds)
        formulas = formulas_book[table.sheet].iter_rows(**bounds)
        headers = list(next(values))
        next(formulas)
        rows = []
        for number, (value_row, formula_row) in enumerate(zip(values, formulas), start=min_row + 1):
            width = max_col - min_col + 1
            value_row = tuple(value_row) + (None,) * (width - len(value_row))
            formula_row = tuple(formula_row) + (None,) * (width - len(formula_row))
            rows.append((number, [
                Cell(value, isinstance(formula, str) and formula.startswith("="))
                for value, formula in zip(value_row, formula_row)
            ]))
        return TableData(table=table, headers=headers, rows=rows)
    finally:
        values_book.close()
        formulas_book.close()


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as file:
        for chunk in iter(lambda: file.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()
