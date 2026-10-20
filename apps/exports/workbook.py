"""The data workbooks (docs/REQUISITOS.md §10), written row by row so that all the history fits.

- Long: one row per measurement, with whether it is suspect and who registered it. One sheet per
  year, so that no sheet passes Excel's limit of a million rows.
- Wide by source: one sheet per source, one row per sampling time, one column per point and
  parameter, like today's sheets. Without suspect values; "<0,1" goes as text.
"""

import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from apps.catalog.models import Source, Variable
from apps.measurements.models import Measurement, TimePrecision
from apps.measurements.values import format_number

LONG = "long"
WIDE = "wide"
SHAPES = [(LONG, "Larga: una fila por medición"), (WIDE, "Ancha: una hoja por fuente, como las hojas de hoy")]
LONG_HEADERS = ["Fecha", "Hora", "Punto", "Parámetro", "Fuente", "Calificador", "Valor", "Unidad", "Dudoso", "Registrada por"]
HOURS = {TimePrecision.AM: "a. m.", TimePrecision.PM: "p. m.", TimePrecision.DAY: ""}
BOLD = Font(bold=True)


def hour_text(local, precision):
    return f"{local:%H:%M}" if precision == TimePrecision.HOUR else HOURS[precision]


def header(sheet, titles):
    cells = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.font = BOLD
        cells.append(cell)
    sheet.append(cells)


def date_cell(sheet, day):
    cell = WriteOnlyCell(sheet, value=day)
    cell.number_format = "yyyy-mm-dd"
    return cell


def sheet_title(text):
    """Excel's rules for sheet names: 31 characters, none of []:*?/\\."""
    return re.sub(r"[\[\]:*?/\\]", "-", text)[:31]


def measurements(start=None, end=None, sources=None):
    """Measurements taken in [start, end), of some sources (None: all)."""
    queryset = Measurement.objects.all()
    if start is not None:
        queryset = queryset.filter(taken_at__gte=start)
    if end is not None:
        queryset = queryset.filter(taken_at__lt=end)
    if sources is not None:
        queryset = queryset.filter(variable__source__in=sources)
    return queryset


def write_long(book, queryset):
    rows = queryset.order_by(
        "taken_at", "variable__point__area__order", "variable__point__order", "variable__point__code",
        "variable__parameter__code", "variable__source__code",
    ).values_list(
        "taken_at", "time_precision", "variable__point__code", "variable__parameter__code", "variable__source__code",
        "qualifier", "value", "variable__parameter__unit", "suspect", "created_by__full_name",
    )
    sheets = {}
    count = 0
    for taken_at, precision, point, parameter, source, qualifier, value, unit, suspect, person in rows.iterator(chunk_size=5000):
        local = timezone.localtime(taken_at)
        sheet = sheets.get(local.year)
        if sheet is None:
            sheet = sheets[local.year] = book.create_sheet(f"Mediciones {local.year}")
            header(sheet, LONG_HEADERS)
        sheet.append([
            date_cell(sheet, local.date()), hour_text(local, precision), point, parameter, source, qualifier, value,
            unit, "sí" if suspect else "", person,
        ])
        count += 1
    if not sheets:
        header(book.create_sheet("Mediciones"), LONG_HEADERS)
    return count


def write_wide(book, queryset):
    queryset = queryset.filter(suspect=False)
    count = 0
    for source in Source.objects.filter(pk__in=queryset.values("variable__source")):
        rows = queryset.filter(variable__source=source)
        variables = list(Variable.objects.filter(pk__in=rows.values("variable")).select_related("point", "parameter"))
        column = {variable.pk: index for index, variable in enumerate(variables, start=2)}
        sheet = book.create_sheet(sheet_title(source.code))
        header(sheet, ["Fecha", "Hora"] + [
            f"{v.point} · {v.parameter.code}" + (f" ({v.parameter.unit})" if v.parameter.unit else "") for v in variables
        ])
        key, line = None, None
        for variable_id, taken_at, precision, value, qualifier in (
            rows.order_by("taken_at", "time_precision")
            .values_list("variable_id", "taken_at", "time_precision", "value", "qualifier")
            .iterator(chunk_size=5000)
        ):
            if (taken_at, precision) != key:
                if line:
                    sheet.append(line)
                key = (taken_at, precision)
                local = timezone.localtime(taken_at)
                line = [date_cell(sheet, local.date()), hour_text(local, precision)] + [None] * len(variables)
            line[column[variable_id]] = f"{qualifier}{format_number(value)}" if qualifier else value
            count += 1
        if line:
            sheet.append(line)
    return count


def write(path, readme, shapes, queryset):
    """Writes the workbook to `path`: a «Léeme» sheet with `readme` (lines), then the shapes asked
    for. Returns how many measurements the first shape holds."""
    book = Workbook(write_only=True)
    about = book.create_sheet("Léeme")
    about.column_dimensions["A"].width = 110
    for line in readme:
        about.append([line])
    counts = []
    if LONG in shapes:
        counts.append(write_long(book, queryset))
    if WIDE in shapes:
        counts.append(write_wide(book, queryset))
    book.save(path)
    return counts[0]


def readme_lines(title, period, shapes):
    lines = [title, period, f"Generado el {timezone.localtime():%Y-%m-%d %H:%M}.", ""]
    if LONG in shapes:
        lines += [
            "Hojas «Mediciones AAAA» (formato largo): una fila por medición, una hoja por año.",
            "  Hora: hh:mm, «a. m.», «p. m.» o vacía cuando la fuente solo da el día.",
            "  Calificador: «<» o «>» delante del valor («<0,1»: por debajo de 0,1).",
            "  Dudoso: valores del histórico fuera del rango físico, pendientes de revisión.",
        ]
    if WIDE in shapes:
        lines += [
            "Hojas por fuente (formato ancho): una fila por fecha y hora, una columna por punto y parámetro.",
            "  Sin los valores dudosos. Los valores con «<» o «>» van como texto.",
        ]
    lines += ["", "Un dato que no hay (ND, FS, SB…) no tiene fila ni celda: queda vacío."]
    return lines
