"""The explorer's Excel export: what is being viewed (docs/REQUISITOS.md §7.2).

- «Mediciones»: one row per measurement, in long format, with who registered it (§10);
- «Estadística»: the table of the page;
- the chart's means, when it shows daily, weekly or monthly means.
Suspect values are left out, as in the page; the statistics sheet says how many.
"""

import io
import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from apps.catalog.models import Variable
from apps.measurements.models import Measurement, TimePrecision

from .stats import aggregate, day_span

HOURS = {TimePrecision.AM: "a. m.", TimePrecision.PM: "p. m.", TimePrecision.DAY: ""}


def header(sheet, *titles):
    cells = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.font = Font(bold=True)
        cells.append(cell)
    sheet.append(cells)


def date_cell(sheet, day):
    cell = WriteOnlyCell(sheet, value=day)
    cell.number_format = "yyyy-mm-dd"
    return cell


def measurements_sheet(book, exploration):
    sheet = book.create_sheet("Mediciones")
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["I"].width = 28
    header(sheet, "Fecha", "Hora", "Punto", "Parámetro", "Fuente", "Calificador", "Valor", "Unidad", "Registrada por")
    start, end = day_span(exploration.first_day, exploration.last_day)
    rows = (
        Measurement.objects.filter(variable__in=exploration.variables, suspect=False, taken_at__gte=start, taken_at__lt=end)
        .order_by("taken_at", "variable__point__area__order", "variable__point__order", "variable__point__code", "variable__source__code")
        .values_list(
            "taken_at", "time_precision", "variable__point__code", "variable__parameter__code", "variable__source__code",
            "qualifier", "value", "variable__parameter__unit", "created_by__full_name",
        )
    )
    for taken_at, precision, point, parameter, source, qualifier, value, unit, person in rows.iterator(chunk_size=5000):
        local = timezone.localtime(taken_at)
        hour = f"{local:%H:%M}" if precision == TimePrecision.HOUR else HOURS[precision]
        sheet.append([date_cell(sheet, local.date()), hour, point, parameter, source, qualifier, value, unit, person])


def statistics_sheet(book, exploration):
    sheet = book.create_sheet("Estadística")
    sheet.column_dimensions["A"].width = 30
    parameter = exploration.parameter
    unit = f" ({parameter.unit})" if parameter.unit else ""
    sheet.append([f"{parameter.name}{unit}"])
    sheet.append(["Desde", date_cell(sheet, exploration.first_day), "Hasta", date_cell(sheet, exploration.last_day)])
    sheet.append([])
    header(
        sheet, "Serie", "Muestras", "Promedio", "Mediana", "Mínimo", "Máximo", "Desviación estándar", "P10", "P90",
        "Con límite vigente", "Fuera de límite", "% fuera de límite", "Dudosos (excluidos)",
    )
    for result in exploration.results:
        d, e = result.description, result.exceedance
        sheet.append([
            f"{result.variable.point} · {result.variable.source}", d.n, d.mean, d.median, d.minimum, d.maximum, d.std,
            d.p10, d.p90, e.checked, e.out, e.percent, result.suspect,
        ])
    sheet.append([])
    for note in (
        "Los valores con «<» o «>» cuentan como su número (el peor caso): «<0,1» como 0,1.",
        "Los valores dudosos (fuera del rango físico, sin revisar) no entran en la estadística.",
        "Desviación estándar muestral (DESVEST.M) y percentiles con interpolación (PERCENTIL.INC).",
        "Cada muestra se compara con los límites vigentes en su fecha; sin límite vigente no se cuenta.",
    ):
        sheet.append([note])


def means_sheet(book, exploration):
    sheet = book.create_sheet(exploration.period_name.capitalize()[:31])
    sheet.column_dimensions["A"].width = 30
    header(sheet, "Serie", "Desde", "Promedio", "Muestras", "Con calificador")
    for variable in exploration.variables:
        for bucket in aggregate(exploration.samples.get(variable.pk, []), exploration.period):
            sheet.append([f"{variable.point} · {variable.source}", date_cell(sheet, bucket.start), bucket.mean, bucket.n, bucket.qualified])


def measurement_count(data):
    """How many measurements an export of the explorer's form would hold."""
    start, end = day_span(data["start"], data["end"])
    variables = Variable.objects.filter(parameter=data["parameter"], point__in=data["points"], source__in=data["sources"])
    return Measurement.objects.filter(variable__in=variables, suspect=False, taken_at__gte=start, taken_at__lt=end).count()


def workbook(exploration):
    book = Workbook(write_only=True)
    statistics_sheet(book, exploration)
    measurements_sheet(book, exploration)
    if exploration.period:
        means_sheet(book, exploration)
    content = io.BytesIO()
    book.save(content)
    content.seek(0)
    return content


def filename(exploration):
    code = re.sub(r"[^\w.-]+", "_", exploration.parameter.code).strip("_") or "parametro"
    return f"explorador_{code}_{exploration.first_day:%Y-%m-%d}_{exploration.last_day:%Y-%m-%d}.xlsx"
