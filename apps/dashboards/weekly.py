"""The charts of «ICAV» (docs/REQUISITOS.md §7.3): the weekly charts of the ICAV report.

A week goes from Monday to Sunday. Each chart is a list of variables configured in the admin; each
series (point + source) has one value per day, the mean of that day's measurements. As everywhere,
a value with a qualifier counts as its number, suspect values are left out and a day without data
leaves a gap.
"""

import datetime
import io
from dataclasses import dataclass, field

from django.utils.text import slugify
from openpyxl import Workbook

from apps.catalog.models import Limit

from .charts import chart_data, epoch, point_color, source_shapes
from .export import date_cell, header
from .models import WeeklyChart
from .stats import LimitBook, Period, aggregate, load_samples

WEEK = datetime.timedelta(days=7)
MARGIN = 6 * 3600  # seconds of axis before Monday and after Sunday


def monday_of(day):
    return day - datetime.timedelta(days=day.weekday())


def last_full_week(today):
    """The Monday of the last week that is over: the one that ended last Sunday."""
    return monday_of(today) - WEEK


@dataclass
class Column:
    """A series with data in the week, as in the table under its chart."""

    variable: object
    name: str
    color: str  # the point's, as on the chart
    shape: str  # the source's
    means: dict  # {day: Bucket}


@dataclass
class Panel:
    item: WeeklyChart
    days: list
    columns: list = field(default_factory=list)
    stale: list = field(default_factory=list)  # variables without data in the week
    chart: dict | None = None

    @property
    def chart_id(self):
        return f"grafica-icav-{self.item.pk}"

    @property
    def image_name(self):
        return f"icav_{self.days[0]:%Y-%m-%d}_{slugify(self.item.title)}.png"

    @property
    def image_title(self):
        """The title drawn on the downloaded image: 'Fenoles (mg/L) · semana del … al …'."""
        unit = f" ({self.chart['unit']})" if self.chart and self.chart["unit"] else ""
        return f"{self.item.title}{unit} · semana del {self.days[0]:%Y-%m-%d} al {self.days[-1]:%Y-%m-%d}"

    @property
    def decimals(self):
        return max((column.variable.parameter.decimals for column in self.columns), default=None)

    def rows(self):
        """(day, [Bucket or None per column]) for each day of the week."""
        return [(day, [column.means.get(day) for column in self.columns]) for day in self.days]


def unique(items):
    """The items in order, each once."""
    return list(dict.fromkeys(items))


def build_panel(item, monday):
    sunday = monday + datetime.timedelta(days=6)
    panel = Panel(item, days=[monday + datetime.timedelta(days=n) for n in range(7)])
    variables = [entry.variable for entry in item.items.all()]
    if not variables:
        return panel
    points = unique(variable.point for variable in variables)
    parameters = unique(variable.parameter for variable in variables)
    limits = Limit.objects.filter(parameter__in=parameters, point__in=points).select_related("point", "parameter")
    books = {parameter.pk: LimitBook([limit for limit in limits if limit.parameter_id == parameter.pk]) for parameter in parameters}
    samples = load_samples(variables, monday, sunday)

    colors = {point.pk: point_color(index) for index, point in enumerate(points)}
    shapes = source_shapes()
    several = len(parameters) > 1
    for variable in variables:
        rows = samples.get(variable.pk)
        if not rows:
            panel.stale.append(variable)
            continue
        name = f"{variable.point} · {variable.source}"
        panel.columns.append(Column(
            variable=variable,
            name=f"{variable.parameter.code} · {name}" if several else name,
            color=colors[variable.point_id],
            shape=shapes[variable.source_id][0],
            means={bucket.start: bucket for bucket in aggregate(rows, Period.DAY)},
        ))

    panel.chart = chart_data(parameters, points, variables, samples, books, monday, sunday, Period.DAY)
    # A day's mean is drawn at its 00:00: the axis goes from Monday's to Sunday's, with a quarter of
    # a day on each side so their markers are not cut, and it is labeled by day ("lun 19").
    panel.chart["start"] = epoch(monday) - MARGIN
    panel.chart["end"] = epoch(sunday) + MARGIN
    panel.chart["days"] = True
    return panel


def build_panels(monday):
    items = WeeklyChart.objects.prefetch_related(
        "items__variable__point", "items__variable__parameter", "items__variable__source"
    )
    return [build_panel(item, monday) for item in items]


# ---------------------------------------------------------------- Excel


def sheet_title(title, taken):
    """An Excel sheet name: at most 31 characters, none of []:*?/\\, and not repeated."""
    clean = " ".join("".join(" " if char in "[]:*?/\\" else char for char in title).split())[:31].strip() or "Gráfica"
    name, n = clean, 2
    while name.lower() in taken:
        suffix = f" ({n})"
        name, n = clean[: 31 - len(suffix)] + suffix, n + 1
    taken.add(name.lower())
    return name


def workbook(panels, monday):
    """One sheet per chart: the days in rows, the series in columns, the daily means."""
    book = Workbook(write_only=True)
    sunday = monday + datetime.timedelta(days=6)
    taken = set()
    for panel in panels:
        sheet = book.create_sheet(sheet_title(panel.item.title, taken))
        sheet.column_dimensions["A"].width = 12
        units = unique(column.variable.parameter.unit for column in panel.columns)
        sheet.append([panel.item.title + (f" ({units[0]})" if len(units) == 1 and units[0] else "")])
        sheet.append(["Semana", date_cell(sheet, monday), "a", date_cell(sheet, sunday)])
        sheet.append([])
        if not panel.columns:
            sheet.append(["Sin datos en la semana."])
            continue
        header(sheet, "Día", *(column.name for column in panel.columns))
        for day, buckets in panel.rows():
            sheet.append([date_cell(sheet, day), *(bucket.mean if bucket else None for bucket in buckets)])
        sheet.append([])
        sheet.append(["Promedio de las mediciones de cada día. Los valores con «<» o «>» cuentan como su número (el peor caso)."])
        qualified = [
            f"{column.name}, {day:%Y-%m-%d}"
            for column in panel.columns
            for day, bucket in sorted(column.means.items())
            if bucket.qualified
        ]
        if qualified:
            sheet.append(["Promedios con valores «<» o «>»: " + "; ".join(qualified) + "."])
        if panel.stale:
            sheet.append(["Sin datos en la semana: " + "; ".join(f"{v.point} · {v.parameter.code} · {v.source}" for v in panel.stale) + "."])
    if not panels:
        book.create_sheet("ICAV").append(["No hay gráficas del ICAV configuradas."])
    content = io.BytesIO()
    book.save(content)
    content.seek(0)
    return content


def filename(monday):
    return f"icav_semana_{monday:%Y-%m-%d}.xlsx"
