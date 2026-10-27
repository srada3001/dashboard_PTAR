"""The ICAV charts: weeks from Monday to Sunday, daily means, configuration and export (docs/REQUISITOS.md §7.3)."""

import datetime
import importlib
import io
import json
import re

import pytest
from django.apps import apps as django_apps
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.catalog.models import Limit, Parameter, SamplingPoint, Source, Variable
from apps.core.models import AuditEntry
from apps.dashboards import weekly
from apps.dashboards.charts import epoch
from apps.dashboards.models import MAX_POINTS, WeeklyChart, WeeklyChartVariable
from apps.measurements.models import TimePrecision

from .test_dashboards import at, measure

pytestmark = pytest.mark.django_db

MONDAY = datetime.date(2026, 10, 19)
SUNDAY = datetime.date(2026, 10, 25)


def day(n):
    """The n-th day of the week (0: Monday)."""
    return MONDAY + datetime.timedelta(days=n)


@pytest.fixture
def labviewer(db):
    return Source.objects.create(code="Labviewer", name="Laboratorio central")


@pytest.fixture
def phenols_lab(outlet, phenols, labviewer):
    return Variable.objects.create(point=outlet, parameter=phenols, source=labviewer)


@pytest.fixture
def chart(phenols_biota, phenols_lab):
    item = WeeklyChart.objects.create(title="Fenoles", order=1)
    WeeklyChartVariable.objects.create(chart=item, variable=phenols_biota, order=1)
    WeeklyChartVariable.objects.create(chart=item, variable=phenols_lab, order=2)
    return item


def only_panel(monday=MONDAY):
    [panel] = weekly.build_panels(monday)
    return panel


class TestWeeks:
    def test_a_week_goes_from_monday_to_sunday(self):
        assert weekly.monday_of(day(3)) == MONDAY
        assert weekly.monday_of(MONDAY) == MONDAY
        assert weekly.monday_of(SUNDAY) == MONDAY

    def test_by_default_the_last_week_that_is_over(self):
        assert weekly.last_full_week(day(7)) == MONDAY  # on the next Monday
        assert weekly.last_full_week(SUNDAY) == MONDAY - weekly.WEEK  # on Sunday the week is not over

    def test_the_page_shows_the_week_of_the_day_chosen(self, engineer_client, chart):
        response = engineer_client.get(reverse("weekly_charts"), {"semana": day(3).isoformat()})
        assert response.context["monday"] == MONDAY and response.context["sunday"] == SUNDAY

    def test_a_wrong_or_future_day_falls_back(self, engineer_client, chart):
        today = timezone.localdate()
        wrong = engineer_client.get(reverse("weekly_charts"), {"semana": "2026-13-40"})
        assert wrong.context["monday"] == weekly.last_full_week(today)
        future = engineer_client.get(reverse("weekly_charts"), {"semana": "2100-01-01"})
        assert future.context["monday"] == weekly.monday_of(today)
        assert future.context["next_monday"] is None


class TestDailyMeans:
    def test_one_value_per_day_the_mean_of_its_measurements(self, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        measure(phenols_biota, admin_user, at(MONDAY, 16), 0.3)
        panel = only_panel()
        means = panel.columns[0].means
        assert means[MONDAY].mean == pytest.approx(0.2) and means[MONDAY].n == 2

    def test_only_the_days_from_monday_to_sunday(self, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY - datetime.timedelta(days=1), 23, 30), 9)  # Sunday before
        measure(phenols_biota, admin_user, at(MONDAY, 0), 1)
        measure(phenols_biota, admin_user, at(SUNDAY, 23, 30), 2)
        measure(phenols_biota, admin_user, at(day(7), 0), 9)  # next Monday
        panel = only_panel()
        assert {d: b.mean for d, b in panel.columns[0].means.items()} == {MONDAY: 1, SUNDAY: 2}

    def test_a_sample_known_only_by_its_day_counts_in_that_day(self, chart, phenols_lab, admin_user):
        measure(phenols_lab, admin_user, at(day(2)), 0.15, precision=TimePrecision.DAY)
        panel = only_panel()
        assert panel.columns[0].variable == phenols_lab
        assert list(panel.columns[0].means) == [day(2)]

    def test_qualified_values_count_as_their_number_and_are_marked(self, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1, qualifier="<")
        measure(phenols_biota, admin_user, at(MONDAY, 16), 0.3)
        bucket = only_panel().columns[0].means[MONDAY]
        assert bucket.mean == pytest.approx(0.2) and bucket.qualified == 1

    def test_suspect_values_are_left_out(self, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        measure(phenols_biota, admin_user, at(MONDAY, 16), 73, suspect=True)
        assert only_panel().columns[0].means[MONDAY].mean == pytest.approx(0.1)

    def test_a_day_without_data_is_a_gap_not_a_value(self, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        measure(phenols_biota, admin_user, at(day(2), 7), 0.2)
        panel = only_panel()
        rows = dict(panel.rows())
        assert rows[day(1)] == [None] and len(rows) == 7
        assert panel.chart["gap"] is not None  # charts.js breaks the line at a missing day

    def test_series_without_data_are_listed_apart(self, chart, phenols_biota, phenols_lab, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        panel = only_panel()
        assert [c.variable for c in panel.columns] == [phenols_biota]
        assert panel.stale == [phenols_lab]

    def test_no_data_at_all_is_no_chart(self, chart):
        panel = only_panel()
        assert panel.columns == [] and len(panel.stale) == 2


class TestChart:
    def test_the_axis_spans_the_week_and_draws_daily_means(self, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        data = only_panel().chart
        assert (data["start"], data["end"], data["period"]) == (epoch(MONDAY) - weekly.MARGIN, epoch(SUNDAY) + weekly.MARGIN, "day")
        assert data["series"][0]["data"][0][:2] == [epoch(MONDAY), 0.1]

    def test_color_is_the_point_and_shape_the_source(self, chart, phenols_biota, phenols_lab, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        measure(phenols_lab, admin_user, at(MONDAY), 0.12, precision=TimePrecision.DAY)
        data = only_panel().chart
        assert [p["code"] for p in data["points"]] == ["BA-4010"]
        assert {s["point"] for s in data["series"]} == {0}
        assert len({s["source"] for s in data["series"]}) == 2

    def test_several_parameters_share_the_chart_when_they_share_the_unit(self, outlet, biota, admin_user):
        nt = Parameter.objects.create(code="N total", name="Nitrógeno total", unit="mg/L")
        nh3 = Parameter.objects.create(code="N-NH3", name="Nitrógeno amoniacal", unit="mg/L")
        item = WeeklyChart.objects.create(title="Nitrógeno")
        for order, parameter in enumerate((nt, nh3)):
            variable = Variable.objects.create(point=outlet, parameter=parameter, source=biota)
            WeeklyChartVariable.objects.create(chart=item, variable=variable, order=order)
            measure(variable, admin_user, at(MONDAY, 7), 50 + order)
        panel = only_panel()
        assert [p["code"] for p in panel.chart["parameters"]] == ["N total", "N-NH3"]
        assert [c.name for c in panel.columns] == ["N total · BA-4010 · BIOTA", "N-NH3 · BA-4010 · BIOTA"]

    def test_the_limit_in_force_is_drawn(self, chart, phenols_biota, outlet, phenols, admin_user):
        Limit.objects.create(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2)
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        [limit] = only_panel().chart["limits"]
        assert limit["values"] == [0.2]


class TestPage:
    def test_chart_and_table_of_the_week(self, engineer_client, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1)
        measure(phenols_biota, admin_user, at(MONDAY, 16), 0.2)
        html = engineer_client.get(reverse("weekly_charts"), {"semana": MONDAY.isoformat()}).content.decode()
        assert "Gráficas del ICAV" in html and "Fenoles" in html
        assert "0,15" in html  # the mean, with a decimal comma
        assert "Sin datos en la semana:" in html and "BA-4010 · Labviewer" in html
        assert 'data-image="icav_2026-10-19_fenoles.png"' in html
        data = json.loads(re.search(r'<script id="grafica-icav-\d+" type="application/json">(.*?)</script>', html).group(1))
        assert data["period"] == "day"

    def test_the_menu_has_the_tab(self, engineer_client, chart):
        assert ">ICAV</a>" in engineer_client.get(reverse("important_variables")).content.decode()

    def test_without_configuration_it_says_where_to_set_it(self, admin_client):
        html = admin_client.get(reverse("weekly_charts")).content.decode()
        assert "No hay gráficas del ICAV configuradas." in html
        assert reverse("admin:dashboards_weeklychart_changelist") in html


class TestExport:
    def test_one_sheet_per_chart_with_the_daily_means(self, engineer_client, chart, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(MONDAY, 7), 0.1, qualifier="<")
        measure(phenols_biota, admin_user, at(MONDAY, 16), 0.3)
        response = engineer_client.get(reverse("weekly_charts_export"), {"semana": day(4).isoformat()})
        assert response["Content-Disposition"].endswith('filename="icav_semana_2026-10-19.xlsx"')
        book = load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        assert book.sheetnames == ["Fenoles"]
        rows = list(book["Fenoles"].iter_rows(values_only=True))
        assert rows[0][0] == "Fenoles (mg/L)"
        assert rows[3][:2] == ("Día", "BA-4010 · BIOTA")
        assert rows[4][0].date() == MONDAY and rows[4][1] == pytest.approx(0.2)
        assert rows[5][1] is None  # Tuesday: no data
        notes = " ".join(str(r[0]) for r in rows[12:] if r and r[0])
        assert "BA-4010 · BIOTA, 2026-10-19" in notes  # the mean with "<"
        assert "Sin datos en la semana: BA-4010 · Fenoles · Labviewer" in notes

    def test_sheet_names_are_valid_and_unique(self):
        taken = set()
        assert weekly.sheet_title("Sulfatos: separadores/PTAR [mg/L]", taken) == "Sulfatos separadores PTAR mg L"
        assert weekly.sheet_title("Fenoles", taken) == "Fenoles"
        assert weekly.sheet_title("fenoles", taken) == "fenoles (2)"


class TestAdmin:
    def post(self, client, variables):
        data = {
            "title": "Prueba", "order": 1,
            "items-TOTAL_FORMS": str(len(variables)), "items-INITIAL_FORMS": "0",
        }
        for n, variable in enumerate(variables):
            data[f"items-{n}-variable"] = variable.pk
            data[f"items-{n}-order"] = n
        return client.post(reverse("admin:dashboards_weeklychart_add"), data)

    def test_a_chart_is_saved_and_audited(self, admin_client, phenols_biota, phenols_lab):
        response = self.post(admin_client, [phenols_biota, phenols_lab])
        assert response.status_code == 302
        assert WeeklyChart.objects.get().items.count() == 2
        entry = AuditEntry.objects.get(object_type="gráfica del ICAV")
        assert entry.changes["Serie de la gráfica «BA-4010 · Fenoles · BIOTA»"] == ["", "agregado"]

    def test_series_of_different_units_are_rejected(self, admin_client, phenols_biota, outlet, biota):
        temperature = Parameter.objects.create(code="Temperatura", name="Temperatura", unit="°C")
        other = Variable.objects.create(point=outlet, parameter=temperature, source=biota)
        response = self.post(admin_client, [phenols_biota, other])
        assert response.status_code == 200
        assert "comparten el eje" in response.content.decode()
        assert not WeeklyChart.objects.exists()

    def test_at_most_eight_points(self, admin_client, phenols, biota):
        variables = [
            Variable.objects.create(point=SamplingPoint.objects.create(code=f"P-{n}", name=f"P-{n}"), parameter=phenols, source=biota)
            for n in range(MAX_POINTS + 1)
        ]
        response = self.post(admin_client, variables)
        assert response.status_code == 200
        assert f"Hasta {MAX_POINTS} puntos" in response.content.decode()
        assert self.post(admin_client, variables[:MAX_POINTS]).status_code == 302


class TestInitialCharts:
    def test_they_are_created_from_the_catalog_once(self, phenols_biota, phenols_lab, outlet, phenols, labviewer):
        clarified = SamplingPoint.objects.create(code="BA-4010 clarificada", name="BA-4010 clarificada")
        lab_clarified = Variable.objects.create(point=clarified, parameter=phenols, source=labviewer)
        migration = importlib.import_module("apps.dashboards.migrations.0004_initial_weekly_charts")
        migration.create_weekly_charts(django_apps, None)
        migration.create_weekly_charts(django_apps, None)
        item = WeeklyChart.objects.get()  # only «Fenoles» has variables in this catalog
        assert item.title == "Fenoles"
        # BA-4010 · Labviewer is not in the chart: the pivot tables use «BA-4010 clarificada».
        assert [entry.variable for entry in item.items.all()] == [phenols_biota, lab_clarified]
