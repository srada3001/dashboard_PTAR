"""The reading views: windows, statistics, limits, gaps and suspect values (docs/REQUISITOS.md §7, §8)."""

import calendar
import datetime
import importlib
import io
import json
import re

import pytest
from django.apps import apps as django_apps
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.catalog.models import Limit, Parameter, SamplingPoint, Variable
from apps.dashboards.admin import ImportantVariableForm
from apps.dashboards.charts import epoch
from apps.dashboards.models import ImportantVariable
from apps.dashboards.stats import LimitBook, Period, Sample, aggregate, describe, load_samples, within
from apps.measurements.models import Measurement, Origin, TimePrecision, period_start

pytestmark = pytest.mark.django_db

DAY = datetime.date(2026, 10, 19)  # a Monday


def at(day, hour=None, minute=0):
    """An aware instant in Bogotá; without an hour, the start of the day."""
    if hour is None:
        return period_start(day, TimePrecision.DAY)
    return timezone.make_aware(datetime.datetime.combine(day, datetime.time(hour, minute)))


def days_before(count, day=DAY):
    return day - datetime.timedelta(days=count)


def measure(variable, user, when, value, qualifier="", precision=TimePrecision.HOUR, suspect=False):
    return Measurement.objects.create(
        variable=variable, taken_at=when, time_precision=precision, value=value, qualifier=qualifier,
        suspect=suspect, origin=Origin.HISTORY, created_by=user,
    )


def sample(value, when=None, qualifier=""):
    return Sample(when or at(DAY, 8), TimePrecision.HOUR, value, qualifier)


@pytest.fixture
def entry(db):
    return SamplingPoint.objects.create(code="BA-4001", name="Entrada a la PTAR")


@pytest.fixture
def phenols_card(phenols, outlet, entry):
    item = ImportantVariable.objects.create(parameter=phenols, order=1)
    item.points.set([entry, outlet])
    return item


def page(client, **params):
    return client.get(reverse("important_variables"), params).content.decode()


class TestDescriptiveStatistics:
    def test_like_excel(self):
        d = describe([sample(v) for v in (1, 2, 3, 4, 10)])
        assert (d.n, d.mean, d.median, d.minimum, d.maximum) == (5, 4, 3, 1, 10)
        assert d.std == pytest.approx(3.5355339)  # STDEV.S
        assert d.p10 == pytest.approx(1.4)  # PERCENTILE.INC
        assert d.p90 == pytest.approx(7.6)

    def test_one_value_has_no_standard_deviation(self):
        d = describe([sample(7.4)])
        assert (d.n, d.mean, d.std, d.p10, d.p90) == (1, 7.4, None, 7.4, 7.4)

    def test_no_values_is_no_data(self):
        d = describe([])
        assert d.n == 0 and d.mean is None and d.maximum is None

    def test_qualified_values_count_as_the_worst_case(self):
        d = describe([sample(0.1, qualifier="<"), sample(0.3), sample(500, qualifier=">")])
        assert (d.minimum, d.maximum) == (0.1, 500)
        assert d.mean == pytest.approx(166.8)


class TestWindows:
    def test_seven_and_thirty_calendar_days_up_to_the_chosen_day(self, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(days_before(6)), 1, precision=TimePrecision.DAY)  # first of the 7
        measure(phenols_biota, admin_user, at(days_before(7), 23, 30), 2)  # the day before: only in the 30
        measure(phenols_biota, admin_user, at(DAY, 23, 30), 3)  # late on the last day
        measure(phenols_biota, admin_user, at(days_before(29)), 4, precision=TimePrecision.AM)  # first of the 30
        measure(phenols_biota, admin_user, at(days_before(30), 12), 5, precision=TimePrecision.PM)  # out
        measure(phenols_biota, admin_user, at(DAY + datetime.timedelta(days=1)), 6)  # after the last day
        loaded = load_samples([phenols_biota], days_before(29), DAY)[phenols_biota.pk]
        assert [s.value for s in loaded] == [4, 2, 1, 3]
        assert [s.value for s in within(loaded, days_before(6), DAY)] == [1, 3]

    def test_suspect_values_are_never_loaded(self, phenols_biota, admin_user):
        measure(phenols_biota, admin_user, at(DAY, 8), 0.2)
        measure(phenols_biota, admin_user, at(DAY, 9), 73, suspect=True)
        assert [s.value for s in load_samples([phenols_biota], DAY, DAY)[phenols_biota.pk]] == [0.2]


class TestLimits:
    def test_each_sample_is_checked_against_the_limit_in_force_on_its_date(self, outlet, phenols):
        old = Limit(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2, valid_to=days_before(5))
        new = Limit(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.5, valid_from=days_before(4))
        book = LimitBook([old, new])
        assert book.breaks(outlet.pk, sample(0.3, at(days_before(6), 8))) is True
        assert book.breaks(outlet.pk, sample(0.3, at(days_before(4), 8))) is False
        assert book.breaks(outlet.pk, sample(0.5, at(DAY, 8))) is False  # the limit itself complies

    def test_a_window_breaks_on_both_sides(self, outlet, ph):
        book = LimitBook([Limit(point=outlet, parameter=ph, kind=Limit.Kind.OPERATING, min_value=6.4, max_value=8.5)])
        assert [book.breaks(outlet.pk, sample(v)) for v in (6.3, 6.4, 7.5, 8.5, 8.6)] == [True, False, False, False, True]

    def test_without_a_limit_nothing_is_checked(self, outlet, entry, phenols):
        book = LimitBook([Limit(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2)])
        assert book.breaks(entry.pk, sample(99)) is None
        result = book.exceedance(entry.pk, [sample(99)])
        assert (result.out, result.checked, result.percent) == (0, 0, None)

    def test_qualified_values_are_checked_as_the_worst_case(self, outlet, phenols):
        book = LimitBook([Limit(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2)])
        result = book.exceedance(outlet.pk, [sample(0.1, qualifier="<"), sample(0.3, qualifier="<"), sample(0.15)])
        assert (result.out, result.checked, result.percent) == (1, 3, pytest.approx(100 / 3))


class TestAggregation:
    def test_means_by_day_week_and_month_in_bogota_time(self):
        samples = [
            sample(1, at(datetime.date(2026, 9, 30), 23)),  # 04:00 UTC on October 1st, still September 30th here
            sample(3, at(datetime.date(2026, 10, 13), 7)),  # Tuesday
            sample(5, at(datetime.date(2026, 10, 18), 16), qualifier="<"),  # Sunday, same week
            sample(7, at(DAY, 7)),  # Monday, next week
        ]
        assert [(b.start, b.mean) for b in aggregate(samples, Period.DAY)][0] == (datetime.date(2026, 9, 30), 1)
        weeks = aggregate(samples, Period.WEEK)
        assert [(b.start, b.mean, b.n, b.qualified) for b in weeks[1:]] == [
            (datetime.date(2026, 10, 12), 4, 2, 1), (DAY, 7, 1, 0),
        ]
        assert [(b.start, b.n) for b in aggregate(samples, Period.MONTH)] == [
            (datetime.date(2026, 9, 1), 1), (datetime.date(2026, 10, 1), 3),
        ]

    def test_days_without_values_have_no_mean(self):
        samples = [sample(1, at(days_before(3), 8)), sample(2, at(DAY, 8))]
        assert [b.start for b in aggregate(samples, Period.DAY)] == [days_before(3), DAY]


class TestChartData:
    def test_times_are_bogotas_wall_clock(self):
        assert epoch(at(DAY, 16)) == calendar.timegm((2026, 10, 19, 16, 0, 0))
        assert epoch(DAY) == calendar.timegm((2026, 10, 19, 0, 0, 0))


class TestImportantVariables:
    def test_last_value_extremes_and_means(self, admin_user, admin_client, phenols_card, phenols_biota):
        measure(phenols_biota, admin_user, at(days_before(20), 7), 0.9)  # only in the 30-day mean
        measure(phenols_biota, admin_user, at(days_before(2), 7), 0.1, qualifier="<")
        measure(phenols_biota, admin_user, at(DAY, 16), 0.3)
        measure(phenols_biota, admin_user, at(DAY, 7), 73, suspect=True)
        html = page(admin_client, end=DAY.isoformat())
        assert "BA-4010" in html and "BIOTA" in html
        assert "0,30" in html and "2026-10-19 16:00" in html  # the last valid value, with the parameter's decimals
        assert "0,10" in html  # minimum of 7 days: "<0,1" counts as 0.1
        assert "0,20" in html  # mean of 7 days
        assert "0,43" in html  # mean of 30 days
        table = html[html.index("<table"):html.index("</table>")]
        chart = json.loads(re.search(r'<script id="grafica-\d+" type="application/json">(.*?)</script>', html).group(1))
        assert "73" not in table  # the suspect value is neither in the table nor in the chart
        assert [row[1] for row in chart["series"][0]["data"]] == [0.9, 0.1, 0.3]

    def test_no_data_in_the_week_is_never_compliance(self, admin_user, admin_client, phenols_card, phenols_biota, outlet, phenols):
        Limit.objects.create(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2)
        measure(phenols_biota, admin_user, at(days_before(10), 7), 0.1)
        html = page(admin_client, end=DAY.isoformat())
        assert "Sin datos" in html
        assert "0 de" not in html

    def test_values_outside_the_limit_are_flagged(self, admin_user, admin_client, phenols_card, phenols_biota, outlet, phenols):
        Limit.objects.create(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2)
        measure(phenols_biota, admin_user, at(days_before(1), 7), 0.1)
        measure(phenols_biota, admin_user, at(DAY, 7), 0.5)
        html = page(admin_client, end=DAY.isoformat())
        assert "Fuera de límite</span>" in html
        assert "▲ 1 de 2" in html
        assert "Límite de vertimiento: ≤ 0,2 mg/L" in html

    def test_only_the_configured_points_are_shown(self, admin_user, admin_client, phenols, biota, outlet, phenols_biota):
        other = SamplingPoint.objects.create(code="SE-3010")
        item = ImportantVariable.objects.create(parameter=phenols)
        item.points.set([outlet])
        measure(phenols_biota, admin_user, at(DAY, 7), 0.1)
        measure(Variable.objects.create(point=other, parameter=phenols, source=biota), admin_user, at(DAY, 7), 4.2)
        html = page(admin_client, end=DAY.isoformat())
        assert "BA-4010" in html and "SE-3010" not in html

    def test_every_source_at_a_point_is_shown(self, admin_user, admin_client, phenols_card, phenols_biota, outlet, phenols, nalco):
        measure(phenols_biota, admin_user, at(DAY, 7), 0.1)
        measure(Variable.objects.create(point=outlet, parameter=phenols, source=nalco), admin_user, at(DAY, 9), 0.12)
        html = page(admin_client, end=DAY.isoformat())
        assert "BIOTA</span>" in html and "Nalco</span>" in html

    def test_series_without_recent_data_go_to_the_footnote(self, admin_user, admin_client, phenols_card, phenols_biota):
        measure(phenols_biota, admin_user, at(days_before(60), 7), 0.1)
        html = page(admin_client, end=DAY.isoformat())
        assert "Sin datos en los últimos 30 días" in html
        assert f"último: {days_before(60).isoformat()} 07:00" in html

    def test_a_future_day_is_refused(self, admin_client, phenols_card):
        tomorrow = timezone.localdate() + datetime.timedelta(days=1)
        assert "No puede ser una fecha futura" in page(admin_client, end=tomorrow.isoformat())

    def test_old_data_is_offered(self, admin_user, admin_client, phenols_card, phenols_biota):
        old = timezone.localdate() - datetime.timedelta(days=12)
        measure(phenols_biota, admin_user, at(old, 7), 0.1)
        html = page(admin_client)
        assert f"el último es del {old.isoformat()}" in html
        assert f"?end={old.isoformat()}&amp;days=30" in html

    def test_without_configuration(self, admin_client, engineer_client):
        assert "No hay variables importantes configuradas" in page(admin_client)
        assert "Pídele al administrador" in page(engineer_client)


class TestImportantVariableAdmin:
    def test_at_most_eight_points(self, phenols):
        points = [SamplingPoint.objects.create(code=f"P-{i}") for i in range(9)]
        form = ImportantVariableForm(data={"parameter": phenols.pk, "order": 1, "points": [p.pk for p in points]})
        assert not form.is_valid()
        assert "Hasta 8 puntos" in str(form.errors["points"])
        form = ImportantVariableForm(data={"parameter": phenols.pk, "order": 1, "points": [p.pk for p in points[:8]]})
        assert form.is_valid()

    def test_changes_are_audited(self, admin_client, phenols, outlet):
        from apps.core.models import AuditEntry

        response = admin_client.post(
            reverse("admin:dashboards_importantvariable_add"), {"parameter": phenols.pk, "order": 1, "points": [outlet.pk]}
        )
        assert response.status_code == 302
        entry = AuditEntry.objects.get(object_type="variable importante")
        assert entry.changes["Puntos"] == ["", "BA-4010"]


class TestInitialImportantVariables:
    def test_they_are_created_from_the_catalog_once(self, phenols, outlet, entry):
        migration = importlib.import_module("apps.dashboards.migrations.0002_initial_important_variables")
        migration.create_important_variables(django_apps, None)
        migration.create_important_variables(django_apps, None)
        item = ImportantVariable.objects.get()
        assert item.parameter == phenols
        assert {p.code for p in item.points.all()} == {"BA-4001", "BA-4010"}  # BA-4008 is not in this catalog


class TestExplorer:
    def query(self, parameter, points, **extra):
        return {"parameter": parameter.pk, "points": [p.pk for p in points], **extra}

    def test_statistics_per_series(self, admin_user, engineer_client, phenols_biota, phenols, outlet):
        for hour, value in enumerate((1, 2, 3, 4, 10), start=7):
            measure(phenols_biota, admin_user, at(DAY, hour), value)
        measure(phenols_biota, admin_user, at(DAY, 20), 73, suspect=True)
        html = engineer_client.get(reverse("explorer"), self.query(phenols, [outlet], start=DAY, end=DAY)).content.decode()
        for shown in ("4,00", "3,00", "1,00", "10,00", "3,54", "1,40", "7,60"):
            assert shown in html
        assert "1 dudoso sin revisar, excluido" in html
        assert "Sin límite" in html

    def test_percent_outside_the_limit(self, admin_user, engineer_client, phenols_biota, phenols, outlet):
        Limit.objects.create(point=outlet, parameter=phenols, kind=Limit.Kind.DISCHARGE, max_value=0.2)
        for hour, value in enumerate((0.1, 0.3, 0.15, 0.25), start=7):
            measure(phenols_biota, admin_user, at(DAY, hour), value)
        html = engineer_client.get(reverse("explorer"), self.query(phenols, [outlet], start=DAY, end=DAY)).content.decode()
        assert "▲ 50,0 %" in html and "2 de 4" in html

    def test_no_samples_is_no_data(self, engineer_client, phenols_biota, phenols, outlet):
        html = engineer_client.get(reverse("explorer"), self.query(phenols, [outlet], start=DAY, end=DAY)).content.decode()
        assert "Sin datos" in html
        assert "0 %" not in html

    def test_points_are_required_and_at_most_eight(self, engineer_client, phenols, biota):
        points = [SamplingPoint.objects.create(code=f"P-{i}") for i in range(9)]
        for point in points:
            Variable.objects.create(point=point, parameter=phenols, source=biota)
        html = engineer_client.get(reverse("explorer"), self.query(phenols, [])).content.decode()
        assert "Elige al menos un punto" in html
        html = engineer_client.get(reverse("explorer"), self.query(phenols, points)).content.decode()
        assert "Elige hasta 8 puntos" in html

    def test_only_the_parameters_points_and_sources_are_offered(self, engineer_client, phenols_biota, phenols, ph, nalco):
        other = SamplingPoint.objects.create(code="SE-3010")
        Variable.objects.create(point=other, parameter=ph, source=nalco)
        html = engineer_client.get(reverse("explorer_options"), {"parameter": phenols.pk}).content.decode()
        assert "BA-4010" in html and "BIOTA" in html
        assert "SE-3010" not in html and "Nalco" not in html
        assert "Elige al menos un punto" not in html  # just chosen: not an error yet

    def test_the_chosen_sources_filter_the_series(self, admin_user, engineer_client, phenols_biota, phenols, outlet, nalco):
        nalco_phenols = Variable.objects.create(point=outlet, parameter=phenols, source=nalco)
        measure(phenols_biota, admin_user, at(DAY, 7), 0.1)
        measure(nalco_phenols, admin_user, at(DAY, 7), 0.2)
        query = self.query(phenols, [outlet], start=DAY, end=DAY, sources=[nalco.pk])
        html = engineer_client.get(reverse("explorer"), query).content.decode()
        assert "Nalco</span>" in html and "BIOTA</span>" not in html

    def test_long_ranges_are_aggregated_automatically(self, engineer_client, phenols_biota, phenols, outlet):
        query = self.query(phenols, [outlet], start=days_before(365), end=DAY)
        html = engineer_client.get(reverse("explorer"), query).content.decode()
        assert "gráfica con promedios diarios (automático)" in html

    def test_export(self, admin_user, engineer_client, phenols_biota, phenols, outlet):
        measure(phenols_biota, admin_user, at(DAY, 7), 0.1, qualifier="<")
        measure(phenols_biota, admin_user, at(DAY), 0.3, precision=TimePrecision.AM)
        measure(phenols_biota, admin_user, at(DAY, 9), 73, suspect=True)
        response = engineer_client.get(reverse("explorer_export"), self.query(phenols, [outlet], start=DAY, end=DAY))
        assert response.status_code == 200
        assert response["Content-Disposition"] == 'attachment; filename="explorador_Fenoles_2026-10-19_2026-10-19.xlsx"'
        book = load_workbook(io.BytesIO(b"".join(response.streaming_content)), read_only=True)
        rows = list(book["Mediciones"].iter_rows(values_only=True))
        assert rows[0][:7] == ("Fecha", "Hora", "Punto", "Parámetro", "Fuente", "Calificador", "Valor")
        assert [(r[1], r[5], r[6]) for r in rows[1:]] == [("a. m.", None, 0.3), ("07:00", "<", 0.1)]
        stats = list(book["Estadística"].iter_rows(values_only=True))
        assert stats[4][:3] == ("BA-4010 · BIOTA", 2, pytest.approx(0.2))
        assert stats[4][-1] == 1  # the suspect value, left out

    def test_an_invalid_export_goes_back_to_the_form(self, engineer_client, phenols):
        response = engineer_client.get(reverse("explorer_export"), {"parameter": phenols.pk})
        assert response.status_code == 302 and response.url.startswith(reverse("explorer"))


class TestCapturerReadsNothing:
    @pytest.mark.parametrize("name", ["explorer_options", "explorer_export"])
    def test_explorer_data_is_forbidden(self, capturer_client, phenols_biota, phenols, outlet, name):
        response = capturer_client.get(reverse(name), {"parameter": phenols.pk, "points": [outlet.pk]})
        assert response.status_code == 403
