"""Historical import (docs/REQUISITOS.md §9), with the real cases of §16: "<10" (qualifier),
"7,,4" and "SB" (NaN), pH 73 (suspect), hours at :59, duplicates, and idempotency."""

import csv
import datetime

import openpyxl
import pytest
from django.core.management import call_command
from django.urls import reverse
from openpyxl.worksheet.table import Table

from apps.catalog.models import Parameter, SamplingPoint, Source, Variable
from apps.core.models import AuditEntry
from apps.imports.importer import ImportFailed, import_workbook
from apps.imports.maps import MapError, read_table_map
from apps.imports.models import ImportBatch
from apps.imports.workbook import find_tables
from apps.measurements.models import Measurement, MeasurementHistory, TimePrecision

pytestmark = pytest.mark.django_db

NOW = datetime.datetime(2026, 10, 18, 12, tzinfo=datetime.UTC)
MAP_HEADER = ["header", "action", "point", "parameter", "source", "time", "frequency", "note"]


def write_workbook(path, sheets):
    """sheets: {(sheet name, table name): [header row, *rows]}"""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for (sheet_name, table_name), rows in sheets.items():
        sheet = book.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)
        last = openpyxl.utils.get_column_letter(len(rows[0]))
        sheet.add_table(Table(displayName=table_name, ref=f"A1:{last}{len(rows)}"))
    book.save(path)
    return path


def write_map(directory, table, rows):
    directory.mkdir(exist_ok=True)
    with open(directory / f"{table}.csv", "w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file, lineterminator="\n")
        writer.writerow(MAP_HEADER)
        writer.writerows(rows)
    return directory


def hourly(day, hour, minute=0, second=0, micro=0):
    return datetime.datetime(2024, 1, day, hour, minute, second, micro)


@pytest.fixture
def maps(tmp_path):
    directory = tmp_path / "maps"
    write_map(directory, "Tabla1", [
        ["Fecha", "date", "", "", "", "", "", ""],
        ["pH BA4001", "import", "BA-4001", "pH", "Nalco", "row", "", ""],
        ["SST\\nIN DNF", "import", "DNF entrada", "SST", "Nalco", "row", "", ""],
        ["% remo", "skip", "", "", "", "", "", "Fórmula."],
    ])
    write_map(directory, "FASE1", [
        ["Fecha", "date", "", "", "", "", "", ""],
        ["BA4010 N-TOTAL", "import", "BA-4010", "N total", "BIOTA", "day", "", ""],
        ["BA4010 FENOL 7 AM", "import", "BA-4010", "Fenoles", "BIOTA", "07:00", "", ""],
        ["SE3010 Fenol pm", "import", "SE-3010", "Fenoles", "Por confirmar", "pm", "", ""],
    ])
    return directory


@pytest.fixture
def table1(tmp_path):
    return write_workbook(tmp_path / "tabla1.xlsx", {("Entrada de Datos", "Tabla1"): [
        ["Fecha", "pH BA4001", "SST\nIN DNF", "% remo", "Columna sin mapeo"],
        [hourly(3, 5, 59, 59, 660000), 7.2, 120, 50, 1],
        [hourly(3, 6, 59, 59, 650000), 73, "SB", 50, 1],  # pH 73: suspect; SB: NaN
        [hourly(3, 8), "7,,4", "=SUM(1,2)", 50, 1],  # typo: NaN; formula: not a value
        [hourly(3, 9), "6.8", " ", None, 1],  # clean number written as text
        [hourly(3, 10), 7.0, 15, None, 1],
        [hourly(3, 10), 7.0, 16, None, 1],  # 10:00 twice: same pH once, different SST none
        [datetime.datetime(1903, 2, 17, 12), 7.1, 17, None, 1],  # impossible year, neighbours don't fix it
        [hourly(4, 1), 7.3, 18, None, 1],
        [hourly(4, 2), datetime.datetime(1900, 5, 2), 19, None, 1],  # a date Excel put in a value cell
    ]})


@pytest.fixture
def fase1(tmp_path):
    return write_workbook(tmp_path / "tabla2.xlsx", {("FASE I ", "FASE1"): [
        ["Fecha", "BA4010 N-TOTAL", "BA4010 FENOL 7 AM", "SE3010 Fenol pm"],
        [datetime.datetime(2022, 11, 17), "<10", 0.21, 1.6],
        [datetime.datetime(2022, 11, 18), 12.5, "NR", None],
    ]})


def run(path, user, maps, **options):
    return import_workbook(path, user, maps_dir=maps, now=NOW, **options)


def values(variable_label):
    point, parameter, source = variable_label.split(" · ")
    return {
        (m.taken_at.astimezone(datetime.timezone(datetime.timedelta(hours=-5))).replace(tzinfo=None), m.qualifier, m.value)
        for m in Measurement.objects.filter(variable__point__code=point, variable__parameter__code=parameter,
                                            variable__source__code=source)
    }


class TestTable1:
    def test_values_hours_and_report(self, table1, maps, admin_user):
        [batch] = run(table1, admin_user, maps)
        assert values("BA-4001 · pH · Nalco") == {
            (hourly(3, 6), "", 7.2), (hourly(3, 7), "", 73.0), (hourly(3, 9), "", 6.8),
            (hourly(3, 10), "", 7.0), (hourly(4, 1), "", 7.3),
        }
        assert values("DNF entrada · SST · Nalco") == {(hourly(3, 6), "", 120.0), (hourly(4, 1), "", 18.0), (hourly(4, 2), "", 19.0)}
        counts, report = batch.counts, batch.report
        assert counts["imported"] == 8
        assert counts["nan"] == 3  # SB, 7,,4 and the date in a value cell
        assert {v["value"] for v in report["nan_values"]} == {"SB", "7,,4", "1900-05-02 00:00"}
        assert counts["formula_cells"] == 1
        assert counts["text_numbers"] == 1
        assert counts["identical_duplicates"] == 1
        assert counts["conflicts"] == 1
        assert report["conflicts"][0]["taken_at"] == "2024-01-03 10:00"
        assert counts["dropped_rows"] == 1
        assert report["dropped_rows"][0]["reason"] == "Fecha imposible."
        assert report["unmapped_headers"] == ["Columna sin mapeo"]

    def test_values_outside_the_physical_range_are_imported_as_suspect(self, table1, maps, admin_user):
        [batch] = run(table1, admin_user, maps)
        suspect = Measurement.objects.get(suspect=True)
        assert (suspect.value, suspect.variable.parameter.code) == (73.0, "pH")
        assert batch.report["suspect_values"][0]["range"] == "0 – 14"

    def test_every_measurement_has_its_creation_in_the_history(self, table1, maps, admin_user):
        [batch] = run(table1, admin_user, maps)
        assert MeasurementHistory.objects.count() == Measurement.objects.count() == 8
        entry = MeasurementHistory.objects.get(new_value=73.0)
        assert (entry.action, entry.user, entry.batch, entry.new_suspect) == ("create", admin_user, batch, True)
        assert all(m.created_by == admin_user and m.origin == "history" for m in Measurement.objects.all())

    def test_running_it_again_changes_nothing(self, table1, maps, admin_user):
        run(table1, admin_user, maps)
        [again] = run(table1, admin_user, maps)
        assert again.counts["imported"] == 0
        assert again.counts["already_present"] == 8
        assert Measurement.objects.count() == 8
        assert Variable.objects.count() == 2

    def test_a_different_value_already_in_the_database_is_reported_not_changed(self, table1, maps, admin_user):
        run(table1, admin_user, maps)
        Measurement.objects.filter(value=7.3).update(value=7.4)
        [again] = run(table1, admin_user, maps)
        assert again.counts["existing_conflicts"] == 1
        assert again.report["existing_conflicts"][0]["existing"] == "7,4"
        assert Measurement.objects.filter(value=7.4).exists()

    def test_a_dry_run_saves_nothing_but_its_report(self, table1, maps, admin_user):
        [batch] = run(table1, admin_user, maps, dry_run=True)
        assert batch.counts["imported"] == 8
        assert not Measurement.objects.exists()
        assert not MeasurementHistory.objects.exists()
        assert not Variable.objects.exists()
        assert not Source.objects.exists()
        assert ImportBatch.objects.get().dry_run


class TestDailyTables:
    def test_qualifiers_fixed_hours_and_periods(self, fase1, maps, admin_user):
        [batch] = run(fase1, admin_user, maps)
        nitrogen = Measurement.objects.get(variable__parameter__code="N total", qualifier="<")
        assert (nitrogen.value, nitrogen.time_precision) == (10.0, TimePrecision.DAY)
        phenol = Measurement.objects.get(variable__point__code="BA-4010", variable__parameter__code="Fenoles")
        assert phenol.time_precision == TimePrecision.HOUR
        assert values("BA-4010 · Fenoles · BIOTA") == {(datetime.datetime(2022, 11, 17, 7), "", 0.21)}
        separator = Measurement.objects.get(variable__point__code="SE-3010")
        assert separator.time_precision == TimePrecision.PM
        assert values("SE-3010 · Fenoles · Por confirmar") == {(datetime.datetime(2022, 11, 17, 12), "", 1.6)}
        assert batch.counts["qualified"] == 1
        assert batch.report["qualified_values"][0]["value"] == "<10"
        assert batch.report["nan_values"][0]["value"] == "NR"

    def test_variables_get_the_frequency_of_their_columns(self, fase1, maps, admin_user):
        run(fase1, admin_user, maps)
        frequencies = {str(v): v.time_resolution for v in Variable.objects.all()}
        assert frequencies == {
            "BA-4010 · N total · BIOTA": "daily",
            "BA-4010 · Fenoles · BIOTA": "twice_daily",
            "SE-3010 · Fenoles · Por confirmar": "twice_daily",
        }


class TestCatalog:
    def test_the_initial_catalog_is_created_with_its_audit(self, fase1, maps, admin_user):
        [batch] = run(fase1, admin_user, maps)
        assert Parameter.objects.get(code="pH").physical_max == 14
        assert SamplingPoint.objects.get(code="BA-4010").aliases.filter(alias="BA4010").exists()
        assert Source.objects.filter(code="Por confirmar").exists()
        assert "Parámetro «pH»" in batch.report["catalog_created"]
        assert AuditEntry.objects.filter(object_type="variable").count() == 3

    def test_existing_entries_are_not_changed(self, fase1, maps, admin_user):
        Parameter.objects.create(code="PH", name="Potencial de hidrógeno", physical_min=2, physical_max=12)
        run(fase1, admin_user, maps)
        ph = Parameter.objects.get(code__iexact="ph")
        assert (ph.name, ph.physical_min, ph.physical_max) == ("Potencial de hidrógeno", 2, 12)


class TestMaps:
    def test_columns_are_found_by_their_header_not_their_position(self, tmp_path, maps, admin_user):
        path = write_workbook(tmp_path / "moved.xlsx", {("Entrada de Datos", "Tabla1"): [
            ["SST\nIN DNF", "Fecha", "pH BA4001"],
            [100, hourly(3, 6), 7.2],
            [101, hourly(3, 7), 7.3],
        ]})
        [batch] = run(path, admin_user, maps)
        assert values("BA-4001 · pH · Nalco") == {(hourly(3, 6), "", 7.2), (hourly(3, 7), "", 7.3)}
        assert batch.report["missing_headers"] == ["% remo"]

    def test_a_column_naming_something_outside_the_catalog_stops_the_import(self, tmp_path, admin_user):
        directory = write_map(tmp_path / "bad", "Tabla1", [
            ["Fecha", "date", "", "", "", "", "", ""],
            ["pH BA4001", "import", "BA-9999", "pH", "Nalco", "row", "", ""],
        ])
        path = write_workbook(tmp_path / "t.xlsx", {("Entrada de Datos", "Tabla1"): [["Fecha", "pH BA4001"], [hourly(3, 6), 7]]})
        with pytest.raises(ImportFailed, match="punto «BA-9999»"):
            run(path, admin_user, directory)
        assert not Measurement.objects.exists()

    def test_a_map_needs_one_date_column_and_known_actions(self, tmp_path):
        directory = write_map(tmp_path / "m1", "T", [["pH", "import", "BA-4001", "pH", "Nalco", "row", "", ""]])
        with pytest.raises(MapError, match="una columna de fecha"):
            read_table_map(directory / "T.csv")
        directory = write_map(tmp_path / "m2", "T", [["Fecha", "date", "", "", "", "", "", ""], ["pH", "copiar", "", "", "", "", "", ""]])
        with pytest.raises(MapError, match="acción «copiar»"):
            read_table_map(directory / "T.csv")

    def test_the_real_maps_are_valid(self):
        from apps.imports.maps import available_maps

        maps = available_maps()
        assert set(maps) == {"Tabla1", "Laboratorio", "FASE1", "FenolAASepSoda4", "FenolAASepSoda48", "BR_4081", "DIAPAC", "Ecodisel"}
        for path in maps.values():
            read_table_map(path)

    def test_tables_are_found_in_the_workbook(self, table1):
        [table] = find_tables(table1).values()
        assert (table.name, table.sheet, table.ref) == ("Tabla1", "Entrada de Datos", "A1:E10")


class TestCommandAndAdmin:
    def test_the_command_prints_a_summary(self, table1, maps, admin_user, capsys, monkeypatch):
        monkeypatch.setattr("apps.imports.management.commands.import_history.import_workbook",
                            lambda path, user, **kw: import_workbook(path, user, maps_dir=maps, now=NOW, **kw))
        call_command("import_history", str(table1), "--dry-run")
        output = capsys.readouterr().out
        assert "EN SECO" in output
        assert "Mediciones importadas: 8" in output
        assert not Measurement.objects.exists()

    def test_the_report_is_shown_in_the_admin(self, table1, maps, admin_user, admin_client):
        [batch] = run(table1, admin_user, maps)
        response = admin_client.get(reverse("admin:imports_importbatch_change", args=[batch.pk]))
        assert response.status_code == 200
        content = response.content.decode()
        assert "Dudosas: fuera del rango físico" in content
        assert "7,,4" in content

    def test_imported_data_is_read_only_in_the_admin(self, table1, maps, admin_user, admin_client):
        run(table1, admin_user, maps)
        measurement = Measurement.objects.first()
        assert admin_client.get(reverse("admin:measurements_measurement_changelist")).status_code == 200
        response = admin_client.post(reverse("admin:measurements_measurement_delete", args=[measurement.pk]), {"post": "yes"})
        assert response.status_code == 403
        assert Measurement.objects.filter(pk=measurement.pk).exists()
