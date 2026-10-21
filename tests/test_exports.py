"""Exports (docs/REQUISITOS.md §10): long and wide workbooks, the weekly file and its downloads."""

import datetime
import io

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.exports import views as export_views
from apps.exports import weekly, workbook
from apps.exports.middleware import WeeklyExportMiddleware
from apps.exports.models import WeeklyExport
from apps.measurements.models import Measurement, Origin, TimePrecision, period_start

pytestmark = pytest.mark.django_db

MONDAY = datetime.date(2026, 10, 19)


def at(day, hour):
    return timezone.make_aware(datetime.datetime.combine(day, datetime.time(hour)))


def measure(variable, user, when, value, precision=TimePrecision.HOUR, **fields):
    return Measurement.objects.create(
        variable=variable, taken_at=when, time_precision=precision, value=value, origin=Origin.HISTORY, created_by=user, **fields,
    )


def rows(book, sheet):
    return list(book[sheet].iter_rows(values_only=True))


@pytest.fixture
def exports_dir(settings, tmp_path):
    settings.EXPORTS_DIR = tmp_path
    settings.EXPORTS_BUCKET = ""
    return tmp_path


@pytest.fixture
def data(admin_user, phenols_biota, outlet, phenols, nalco):
    from apps.catalog.models import Variable

    nalco_phenols = Variable.objects.create(point=outlet, parameter=phenols, source=nalco)
    measure(phenols_biota, admin_user, at(datetime.date(2025, 12, 31), 7), 0.1, qualifier="<")
    measure(phenols_biota, admin_user, at(datetime.date(2026, 10, 18), 16), 0.3)
    measure(phenols_biota, admin_user, period_start(datetime.date(2026, 10, 17), TimePrecision.DAY), 0.2, precision=TimePrecision.DAY)
    measure(nalco_phenols, admin_user, at(datetime.date(2026, 10, 18), 16), 0.25)
    measure(nalco_phenols, admin_user, at(datetime.date(2026, 10, 18), 17), 99, suspect=True)
    measure(phenols_biota, admin_user, at(MONDAY, 7), 0.4)  # after the weekly cut


def written(shapes, queryset):
    content = io.BytesIO()
    count = workbook.write(content, ["Léeme"], shapes, queryset)
    return count, load_workbook(io.BytesIO(content.getvalue()), read_only=True)


class TestWorkbooks:
    def test_long_one_row_per_measurement_and_one_sheet_per_year(self, data):
        count, book = written([workbook.LONG], workbook.measurements(end=weekly.week_cut(MONDAY)))
        assert count == 5
        assert book.sheetnames == ["Léeme", "Mediciones 2025", "Mediciones 2026"]
        assert rows(book, "Mediciones 2025")[1][:7] == (datetime.datetime(2025, 12, 31), "07:00", "BA-4010", "Fenoles", "BIOTA", "<", 0.1)
        last_year = rows(book, "Mediciones 2026")
        assert [r[1] for r in last_year[1:]] == [None, "16:00", "16:00", "17:00"]  # the day only (no hour), then hours
        assert last_year[-1][8] == "sí"  # suspect values are in the long format, marked
        assert last_year[-1][9] == "Administrador"

    def test_wide_one_sheet_per_source_without_suspect_values(self, data):
        count, book = written([workbook.WIDE], workbook.measurements(end=weekly.week_cut(MONDAY)))
        assert count == 4
        assert book.sheetnames == ["Léeme", "BIOTA", "Nalco"]
        biota = rows(book, "BIOTA")
        assert biota[0] == ("Fecha", "Hora", "BA-4010 · Fenoles (mg/L)")
        assert [r[1:] for r in biota[1:]] == [("07:00", "<0,1"), (None, 0.2), ("16:00", 0.3)]
        assert rows(book, "Nalco")[1:] == [(datetime.datetime(2026, 10, 18), "16:00", 0.25)]


class TestWeekly:
    def test_due_from_monday(self):
        assert weekly.due_week(at(MONDAY, 0)) == MONDAY
        assert weekly.due_week(at(MONDAY + datetime.timedelta(days=6), 23)) == MONDAY
        assert weekly.due_week(at(MONDAY - datetime.timedelta(days=1), 23)) == MONDAY - datetime.timedelta(days=7)

    def test_prepared_once_with_the_data_up_to_sunday(self, data, exports_dir):
        export = weekly.ensure_current(at(MONDAY, 0) + datetime.timedelta(minutes=3), background=False)
        assert export.status == WeeklyExport.Status.DONE
        assert export.file_name == "datos_PTAR_hasta_2026-10-18.xlsx" and export.measurements == 5
        book = load_workbook(exports_dir / export.file_name, read_only=True)
        assert book.sheetnames == ["Léeme", "Mediciones 2025", "Mediciones 2026", "BIOTA", "Nalco"]
        assert weekly.ensure_current(at(MONDAY, 9), background=False) is None

    def test_someone_preparing_it_is_left_alone_unless_it_died(self, exports_dir):
        now = at(MONDAY, 1)
        assert weekly.claim(MONDAY, now) is not None
        assert weekly.claim(MONDAY, now + datetime.timedelta(minutes=10)) is None
        assert weekly.claim(MONDAY, now + datetime.timedelta(minutes=31)) is not None

    def test_a_failure_is_kept_and_tried_again_an_hour_later(self, data, exports_dir, monkeypatch):
        def broken(*args):
            raise OSError("sin espacio")

        monkeypatch.setattr(weekly.storage().__class__, "save", broken)
        export = weekly.ensure_current(at(MONDAY, 0), background=False)
        assert export.status == WeeklyExport.Status.FAILED and "sin espacio" in export.error
        export.finished_at = timezone.now() - datetime.timedelta(minutes=30)
        export.save()
        assert weekly.claim(MONDAY) is None
        export.finished_at = timezone.now() - datetime.timedelta(minutes=61)
        export.save()
        assert weekly.claim(MONDAY) is not None

    def test_the_middleware_looks_every_few_minutes(self, settings, monkeypatch, rf):
        settings.WEEKLY_EXPORT_AUTO = True
        calls = []
        monkeypatch.setattr(weekly, "ensure_current", lambda: calls.append(1))
        middleware = WeeklyExportMiddleware(lambda request: "respuesta")
        assert middleware(rf.get("/")) == "respuesta"
        middleware(rf.get("/"))
        assert calls == [1]
        middleware.next_look = 0
        middleware(rf.get("/"))
        assert calls == [1, 1]


class TestPages:
    def test_readers_only(self, engineer_client, capturer_client):
        assert engineer_client.get(reverse("exports")).status_code == 200
        for name in ("exports", "exports_download", "exports_weekly"):
            assert capturer_client.get(reverse(name)).status_code == 403

    def test_download_now(self, data, engineer_client, biota):
        params = {"start": "2026-10-17", "end": "2026-10-18", "shape": "long", "sources": [biota.pk]}
        response = engineer_client.get(reverse("exports_download"), params)
        assert response.status_code == 200
        assert response["Content-Disposition"] == 'attachment; filename="datos_PTAR_largo_2026-10-17_2026-10-18.xlsx"'
        book = load_workbook(io.BytesIO(b"".join(response.streaming_content)), read_only=True)
        assert [r[6] for r in rows(book, "Mediciones 2026")[1:]] == [0.2, 0.3]

    def test_too_much_at_once_points_to_the_weekly_file(self, data, engineer_client, monkeypatch):
        monkeypatch.setattr(export_views, "MAX_AT_ONCE", 2)
        response = engineer_client.get(reverse("exports_download"), {"start": "2025-01-01", "end": "2026-10-19", "shape": "wide"})
        assert response.status_code == 400
        assert "Son 6 mediciones; al momento se descargan hasta 2. Acorta las fechas, elige menos fuentes" in response.content.decode()

    def test_weekly_files_are_listed_and_downloaded(self, data, exports_dir, engineer_client):
        export = weekly.ensure_current(at(MONDAY, 0), background=False)
        page = engineer_client.get(reverse("exports")).content.decode()
        assert "Domingo 2026-10-18" in page and "Descargar" in page
        response = engineer_client.get(reverse("exports_weekly_download", args=[export.pk]))
        assert response["Content-Disposition"] == 'attachment; filename="datos_PTAR_hasta_2026-10-18.xlsx"'
        content = b"".join(response.streaming_content)
        assert content == (exports_dir / export.file_name).read_bytes()

    def test_a_file_being_prepared_is_refreshed(self, engineer_client):
        WeeklyExport.objects.create(week=MONDAY, status=WeeklyExport.Status.RUNNING, started_at=timezone.now())
        page = engineer_client.get(reverse("exports_weekly")).content.decode()
        assert "Preparándose" in page and 'hx-trigger="every 5s"' in page
        assert engineer_client.get(reverse("exports_weekly_download", args=[1])).status_code == 404


class TestExplorerExport:
    def test_too_much_at_once_goes_back_with_a_message(self, data, engineer_client, phenols, outlet, monkeypatch):
        from apps.dashboards import views as dashboard_views

        monkeypatch.setattr(dashboard_views, "MAX_AT_ONCE", 1)
        query = {"p1-parameter": phenols.pk, "p1-points": [outlet.pk], "start": "2025-01-01", "end": "2026-10-19"}
        response = engineer_client.get(reverse("explorer_export"), query, follow=True)
        assert "archivo semanal" in response.content.decode()
